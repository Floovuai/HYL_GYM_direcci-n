from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


OUT_DIR = Path(r"C:\Users\chval\OneDrive\Escritorio\PLATAFORMA GESTION\outputs\estacionalidad_consolidado")
INPUT = OUT_DIR / "estacionalidad_data.json"
OUTPUT = OUT_DIR / "Historico Estacionalidad HYL.xlsx"

BLUE = "1F4E79"
DARK_BLUE = "17365D"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F3F6F8"
BORDER = Side(style="thin", color="D9E2F3")


def value(row, key):
    val = row.get(key)
    return "" if val is None else val


def add_table(ws, name, rows, columns):
    ws.append([label for _, label, _ in columns])
    for row in rows:
        ws.append([value(row, key) for key, _, _ in columns])

    max_row = max(ws.max_row, 2)
    max_col = len(columns)
    ref = f"A1:{get_column_letter(max_col)}{max_row}"

    table = Table(displayName=name, ref=ref)
    style = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    table.tableStyleInfo = style
    ws.add_table(table)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ref
    ws.sheet_view.showGridLines = False

    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")

    for col_idx, (_, _, number_format) in enumerate(columns, start=1):
        letter = get_column_letter(col_idx)
        width = 12
        for cell in ws[letter]:
            text = str(cell.value or "")
            width = min(max(width, len(text) + 2), 55)
            if cell.row > 1 and number_format:
                cell.number_format = number_format
        ws.column_dimensions[letter].width = width

    for row in ws.iter_rows():
        for cell in row:
            cell.border = Border(bottom=BORDER)
            cell.alignment = Alignment(vertical="top", wrap_text=False)


def build_summary(wb, data):
    ws = wb.active
    ws.title = "Resumen"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:H1")
    ws["A1"] = "Historico de estacionalidad H&L GYM"
    ws["A1"].fill = PatternFill("solid", fgColor=DARK_BLUE)
    ws["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center")

    ws["A3"] = "Metrica"
    ws["B3"] = "Valor"
    for col in ("A", "B"):
        ws[f"{col}3"].fill = PatternFill("solid", fgColor=BLUE)
        ws[f"{col}3"].font = Font(color="FFFFFF", bold=True)
    for index, item in enumerate(data["resumen"], start=4):
        ws[f"A{index}"] = item["metrica"]
        ws[f"B{index}"] = item["valor"]

    for row in ws["A3:B9"]:
        for cell in row:
            cell.border = Border(bottom=BORDER, top=BORDER, left=BORDER, right=BORDER)

    notes = [
        "Uso recomendado",
        "Historico gerencial agregado: metas, ejecutado, ranking y usuarios activos.",
        "No reemplaza un historico transaccional de ventas por cliente/factura.",
        "Los nombres fueron normalizados, pero se conserva el valor original y archivo fuente.",
        "Los temporales de Excel ~$ fueron ignorados.",
        "En usuarios activos se estima la fecha usando el mes del archivo cuando la etiqueta del dia no es confiable.",
    ]
    ws.merge_cells("D3:H3")
    ws["D3"] = notes[0]
    ws["D3"].fill = PatternFill("solid", fgColor=BLUE)
    ws["D3"].font = Font(color="FFFFFF", bold=True)
    for idx, note in enumerate(notes[1:], start=4):
        ws.merge_cells(start_row=idx, start_column=4, end_row=idx, end_column=8)
        cell = ws.cell(idx, 4, note)
        cell.fill = PatternFill("solid", fgColor=LIGHT_GRAY)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    widths = {"A": 28, "B": 22, "D": 28, "E": 18, "F": 18, "G": 18, "H": 18}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in range(4, 9):
        ws.row_dimensions[row].height = 34


def main():
    data = json.loads(INPUT.read_text(encoding="utf-8"))
    wb = Workbook()
    build_summary(wb, data)

    cobertura_cols = [
        ("periodo", "Periodo", None),
        ("anio", "Anio", None),
        ("mes_num", "Mes Num", None),
        ("mes", "Mes", None),
        ("archivo", "Archivo", None),
        ("hojas", "Hojas", None),
        ("tiene_dia_a_dia", "Tiene Dia a Dia", None),
        ("tiene_ranking_sede", "Tiene Ranking Sede", None),
        ("registros_ranking_sede", "Reg Ranking Sede", None),
        ("tiene_ranking_asesores", "Tiene Ranking Asesores", None),
        ("registros_ranking_asesores", "Reg Ranking Asesores", None),
        ("tiene_usuarios_activos", "Tiene Usuarios Activos", None),
        ("registros_usuarios_activos", "Reg Usuarios Activos", None),
        ("fuente_archivo", "Fuente Archivo", None),
    ]
    ws = wb.create_sheet("Cobertura")
    add_table(ws, "tblCobertura", data["cobertura"], cobertura_cols)

    rank_cols = [
        ("periodo", "Periodo", None),
        ("anio", "Anio", None),
        ("mes_num", "Mes Num", None),
        ("mes", "Mes", None),
        ("entidad_normalizada", "Entidad Normalizada", None),
        ("entidad_original", "Entidad Original", None),
        ("puesto", "Puesto", None),
        ("ejecutado_total", "Ejecutado Total", "#,##0"),
        ("presupuesto", "Presupuesto", "#,##0"),
        ("meta_1", "Meta 1", "#,##0"),
        ("meta_1_5", "Meta 1.5", "#,##0"),
        ("meta_2", "Meta 2", "#,##0"),
        ("meta_2_5", "Meta 2.5", "#,##0"),
        ("meta_3", "Meta 3", "#,##0"),
        ("meta_4", "Meta 4", "#,##0"),
        ("porcentaje_cumplimiento", "% Cumplimiento", "0.0%"),
        ("porcentaje_a_hoy", "% A Hoy", "0.0%"),
        ("diferencia_porcentaje", "Diferencia %", "0.0%"),
        ("valor_a_hoy", "$ A Hoy", "#,##0"),
        ("hoja", "Hoja", None),
        ("fila_origen", "Fila Origen", None),
        ("formula_ejecutado", "Formula Ejecutado", None),
        ("fuente_archivo", "Fuente Archivo", None),
    ]
    ws = wb.create_sheet("Ranking Sedes")
    add_table(ws, "tblRankingSedes", data["ranking_sedes"], rank_cols)
    ws = wb.create_sheet("Ranking Asesores")
    add_table(ws, "tblRankingAsesores", data["ranking_asesores"], rank_cols)

    activos_cols = [
        ("periodo", "Periodo", None),
        ("anio", "Anio", None),
        ("mes_num", "Mes Num", None),
        ("mes", "Mes", None),
        ("fecha_estimada", "Fecha Estimada", "yyyy-mm-dd"),
        ("dia", "Dia", None),
        ("fecha_original", "Fecha Original", None),
        ("sede_normalizada", "Sede Normalizada", None),
        ("sede_original", "Sede Original", None),
        ("usuarios_activos", "Usuarios Activos", "#,##0"),
        ("hoja", "Hoja", None),
        ("fila_origen", "Fila Origen", None),
        ("formula_valor", "Formula Valor", None),
        ("fuente_archivo", "Fuente Archivo", None),
    ]
    ws = wb.create_sheet("Usuarios Activos")
    add_table(ws, "tblUsuariosActivos", data["usuarios_activos"], activos_cols)

    sedes_mes = {}
    for row in data["ranking_sedes"]:
        key = (row["periodo"], row["anio"], row["mes_num"], row["entidad_normalizada"])
        item = sedes_mes.setdefault(
            key,
            {
                "periodo": row["periodo"],
                "anio": row["anio"],
                "mes_num": row["mes_num"],
                "sede": row["entidad_normalizada"],
                "ejecutado_total": 0,
                "registros": 0,
            },
        )
        item["ejecutado_total"] += row.get("ejecutado_total") or 0
        item["registros"] += 1
    sedes_mes_rows = sorted(sedes_mes.values(), key=lambda item: (item["periodo"], item["sede"]))
    sedes_mes_cols = [
        ("periodo", "Periodo", None),
        ("anio", "Anio", None),
        ("mes_num", "Mes Num", None),
        ("sede", "Sede", None),
        ("ejecutado_total", "Ejecutado Total", "#,##0"),
        ("registros", "Registros", None),
    ]
    ws = wb.create_sheet("Sedes Mes")
    add_table(ws, "tblSedesMes", sedes_mes_rows, sedes_mes_cols)

    wb.save(OUTPUT)
    print(f"OK {OUTPUT}")
    print(f"Cobertura: {len(data['cobertura'])}")
    print(f"Ranking sedes: {len(data['ranking_sedes'])}")
    print(f"Ranking asesores: {len(data['ranking_asesores'])}")
    print(f"Usuarios activos: {len(data['usuarios_activos'])}")


if __name__ == "__main__":
    main()
