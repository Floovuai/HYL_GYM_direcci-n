from __future__ import annotations

import json
import re
from datetime import date
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"C:\Users\chval\OneDrive\Escritorio\HYL GESTION COMERCIAL\HYL GESTION ANTERIOR\ESTACIONALIDAD")
OUTPUT = Path(r"C:\Users\chval\OneDrive\Escritorio\PLATAFORMA GESTION\src\server\data\seasonalityDaily.json")

MONTHS = [
    "ENERO",
    "FEBRERO",
    "MARZO",
    "ABRIL",
    "MAYO",
    "JUNIO",
    "JULIO",
    "AGOSTO",
    "SEPTIEMBRE",
    "OCTUBRE",
    "NOVIEMBRE",
    "DICIEMBRE",
]


def clean(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().upper()


def period_from_path(path: Path) -> tuple[int, int]:
    year = int(path.parent.name)
    name = clean(path.name)
    for index, month in enumerate(MONTHS, start=1):
        if month in name:
            return year, index
    raise ValueError(f"No se pudo detectar mes en {path}")


def dia_a_dia_sheet(workbook):
    for sheet in workbook.worksheets:
        if "DIA A DIA" in clean(sheet.title):
            return sheet
    return None


def numeric(value) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    return None


def day_number(value) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        day = int(value)
        return day if 1 <= day <= 31 else None
    match = re.search(r"\d{1,2}", clean(value))
    if not match:
        return None
    day = int(match.group(0))
    return day if 1 <= day <= 31 else None


def header_groups(sheet, row: int) -> list[tuple[int, int]]:
    groups = []
    for col in range(1, sheet.max_column + 1):
        if clean(sheet.cell(row, col).value) != "DIA":
            continue
        for offset in range(1, 4):
            exec_col = col + offset
            if exec_col <= sheet.max_column and clean(sheet.cell(row, exec_col).value) == "EJECUTADO":
                groups.append((col, exec_col))
                break
    return groups


def extract_daily_from_file(path: Path) -> list[dict]:
    year, month = period_from_path(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = dia_a_dia_sheet(workbook)
    if sheet is None:
        return []

    candidates: dict[int, list[dict]] = {}
    for header_row in range(1, sheet.max_row + 1):
        groups = header_groups(sheet, header_row)
        if not groups:
            continue
        for row in range(header_row + 1, sheet.max_row + 1):
            first_label = clean(sheet.cell(row, groups[0][0]).value)
            if "TOTAL" in first_label:
                break
            values = []
            days = []
            for day_col, exec_col in groups:
                day = day_number(sheet.cell(row, day_col).value)
                value = numeric(sheet.cell(row, exec_col).value)
                if day is None or value is None:
                    continue
                days.append(day)
                values.append(value)
            if not days:
                continue
            day = max(set(days), key=days.count)
            if day < 1 or day > 31:
                continue
            try:
                iso_date = date(year, month, day).isoformat()
            except ValueError:
                continue
            record = {
                "year": year,
                "month": month,
                "day": day,
                "date": iso_date,
                "sales": round(sum(values)),
                "sourceFile": path.name,
                "sheet": sheet.title,
                "headerRow": header_row,
                "groupCount": len(values),
            }
            candidates.setdefault(day, []).append(record)

    records = []
    for day, day_records in candidates.items():
        # Some workbooks repeat the same total in advisor and consolidated blocks.
        # Prefer the widest complete block, then the highest sales value.
        records.append(
            max(
                day_records,
                key=lambda item: (
                    item["groupCount"],
                    item["sales"],
                    item["headerRow"],
                ),
            )
        )
    return sorted(records, key=lambda item: item["day"])


def main():
    records = []
    sources = []
    for path in sorted(ROOT.rglob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        try:
            file_records = extract_daily_from_file(path)
        except Exception as error:
            print(f"WARN {path}: {error}")
            continue
        if file_records:
            records.extend(file_records)
        sources.append({"file": path.name, "path": str(path), "dailyRecords": len(file_records)})

    records.sort(key=lambda item: (item["year"], item["month"], item["day"], item["sourceFile"]))
    payload = {
        "generatedAt": date.today().isoformat(),
        "sourceRoot": str(ROOT),
        "note": "Ventas historicas diarias extraidas de todos los bloques DIA/EJECUTADO en hojas DIA A DIA de libros de estacionalidad. No reemplaza ventas transaccionales.",
        "records": records,
        "sources": sources,
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"OK {OUTPUT}")
    print(f"records: {len(records)}")


if __name__ == "__main__":
    main()
