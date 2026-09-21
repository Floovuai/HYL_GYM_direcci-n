from __future__ import annotations

import json
import re
import unicodedata
from datetime import date
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"C:\Users\chval\OneDrive\Escritorio\HYL GESTION COMERCIAL\HYL GESTION ANTERIOR\ESTACIONALIDAD")
OUT_DIR = Path(r"C:\Users\chval\OneDrive\Escritorio\PLATAFORMA GESTION\outputs\estacionalidad_consolidado")
OUT_FILE = OUT_DIR / "estacionalidad_data.json"

MONTHS = [
    "ENERO",
    "FEBRERO",
    "MARZO",
    "ABRIL",
    "MAYO",
    "JUNIO",
    "JULIO",
    "AGOSTO",
    "SEPTIEMBRE",
    "OCTUBRE",
    "NOVIEMBRE",
    "DICIEMBRE",
]

SEDE_ALIASES = {
    "CALLE 109": "109",
    "109": "109",
    "SEDE 109": "109",
    "CALLE 162": "162",
    "162": "162",
    "SEDE 162": "162",
    "BUENOS AIRES": "BUENOS AIRES",
    "MODELIA": "MODELIA",
    "SANTA MATILDE": "SANTA MATILDE",
    "ONLINE": "ONLINE",
    "VILLAVO": "VILLAVICENCIO",
    "VILLAVICENCIO": "VILLAVICENCIO",
    "PRADO": "PRADO VERANIEGO",
    "PRADO VERANIEGO": "PRADO VERANIEGO",
    "FIT SALE": "FIT SALE",
    "CORPORATIVOS": "CORPORATIVOS",
}


def clean_text(value):
    if value is None:
        return ""
    text = str(value).replace("\n", " ").strip()
    text = re.sub(r"\s+", " ", text)
    return text


def norm_key(value):
    text = clean_text(value).upper()
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    text = text.replace("%", "")
    text = text.replace("$", "")
    text = text.replace(",", ".")
    text = re.sub(r"[^A-Z0-9.]+", "_", text).strip("_")
    return text


def normalize_sede(value):
    text = clean_text(value).upper()
    text = text.replace("SEDE ", "").strip()
    text = re.sub(r"\s+", " ", text)
    return SEDE_ALIASES.get(text, text)


def normalize_person(value):
    return clean_text(value).upper()


def period_from_path(path: Path):
    year = int(path.parent.name)
    upper = path.name.upper()
    month = next((i + 1 for i, month_name in enumerate(MONTHS) if month_name in upper), None)
    if not month:
        raise ValueError(f"No se pudo detectar mes en {path}")
    return year, month, MONTHS[month - 1], f"{year}-{month:02d}"


def get_sheet(wb, needle):
    needle = needle.upper()
    for name in wb.sheetnames:
        if needle in name.upper():
            return name
    return None


def is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def formula_at(ws_formula, row, col):
    value = ws_formula.cell(row, col).value
    return value if isinstance(value, str) and value.startswith("=") else ""


def find_header(ws, required_tokens):
    max_rows = min(ws.max_row, 35)
    for row in range(1, max_rows + 1):
        values = [clean_text(ws.cell(row, col).value).upper() for col in range(1, ws.max_column + 1)]
        found = {}
        for token in required_tokens:
            for idx, value in enumerate(values, start=1):
                if token in value:
                    found[token] = idx
                    break
        if len(found) == len(required_tokens):
            return row, found
    return None, {}


def safe_number(value):
    if is_number(value):
        return float(value)
    return None


def header_map(header):
    key = norm_key(header)
    if key == "PUESTO":
        return "puesto"
    if "ASESOR" in key:
        return "entidad_original"
    if key == "SEDE":
        return "entidad_original"
    if "EJECUTADO" in key and "TOTAL" in key:
        return "ejecutado_total"
    if key == "EJECUTADO":
        return "ejecutado_total"
    if "PRESUPUESTO" in key:
        return "presupuesto"
    if key in {"META1", "META_1", "META_1.0", "META_1_"} or key == "META_1":
        return "meta_1"
    if key in {"META2", "META_2"}:
        return "meta_2"
    if key in {"META3", "META_3"}:
        return "meta_3"
    if key in {"META4", "META_4"}:
        return "meta_4"
    if key in {"META_1.5", "META_1_5"}:
        return "meta_1_5"
    if key in {"META_2.5", "META_2_5"}:
        return "meta_2_5"
    if "CUMPLIMIENTO" in key and "HOY" not in key:
        return "porcentaje_cumplimiento"
    if "A_HOY" in key and "DIFERENCIA" not in key:
        return "porcentaje_a_hoy"
    if "DIFERENCIA" in key:
        return "diferencia_porcentaje"
    if key in {"A_HOY", "VALOR_A_HOY"}:
        return "valor_a_hoy"
    if "BONO" in key:
        return None
    return None


def extract_rank(wb_values, wb_formula, sheet_name, label_token, tipo, meta):
    if not sheet_name:
        return []
    ws = wb_values[sheet_name]
    wsf = wb_formula[sheet_name]
    header_row, idx = find_header(ws, ["PUESTO", label_token, "EJECUTADO"])
    if not header_row:
        return []

    headers = [clean_text(ws.cell(header_row, col).value) for col in range(1, ws.max_column + 1)]
    label_col = idx[label_token]
    rows = []
    blank_run = 0
    for row in range(header_row + 1, ws.max_row + 1):
        label = clean_text(ws.cell(row, label_col).value)
        if not label:
            blank_run += 1
            if blank_run >= 8:
                break
            continue
        blank_run = 0
        if label.upper() in {label_token, "TOTAL"} or "TOTAL" in label.upper():
            continue

        record = {
            **meta,
            "tipo": tipo,
            "hoja": sheet_name,
            "fila_origen": row,
            "entidad_original": label,
            "entidad_normalizada": normalize_sede(label) if tipo == "SEDE" else normalize_person(label),
            "puesto": None,
            "ejecutado_total": None,
            "presupuesto": None,
            "meta_1": None,
            "meta_1_5": None,
            "meta_2": None,
            "meta_2_5": None,
            "meta_3": None,
            "meta_4": None,
            "porcentaje_cumplimiento": None,
            "porcentaje_a_hoy": None,
            "diferencia_porcentaje": None,
            "valor_a_hoy": None,
            "formula_ejecutado": "",
        }

        has_value = False
        for col, header in enumerate(headers, start=1):
            mapped = header_map(header)
            if not mapped:
                continue
            value = ws.cell(row, col).value
            if mapped == "entidad_original":
                continue
            if mapped == "puesto":
                record[mapped] = int(value) if is_number(value) else clean_text(value) or None
            else:
                record[mapped] = safe_number(value)
            if mapped == "ejecutado_total":
                record["formula_ejecutado"] = formula_at(wsf, row, col)
            if value not in (None, ""):
                has_value = True

        if has_value and (record["ejecutado_total"] is not None or record["puesto"] is not None):
            rows.append(record)
    return rows


def parse_day(label):
    if isinstance(label, date):
        return label.day
    match = re.search(r"\d{1,2}", clean_text(label))
    if match:
        day = int(match.group(0))
        if 1 <= day <= 31:
            return day
    return None


def extract_activos(wb_values, wb_formula, sheet_name, meta):
    if not sheet_name:
        return []
    ws = wb_values[sheet_name]
    header_row, idx = find_header(ws, ["DIA"])
    if not header_row:
        return []
    dia_col = idx["DIA"]
    headers = [clean_text(ws.cell(header_row, col).value) for col in range(1, ws.max_column + 1)]
    sede_cols = []
    for col in range(dia_col + 1, ws.max_column + 1):
        header = headers[col - 1]
        if not header:
            continue
        header_upper = header.upper()
        if "TOTAL" in header_upper:
            continue
        if "DIA" == header_upper:
            continue
        sede_cols.append((col, header))

    records = []
    for row in range(header_row + 1, ws.max_row + 1):
        label = ws.cell(row, dia_col).value
        day = parse_day(label)
        if not day:
            continue
        try:
            fecha = date(meta["anio"], meta["mes_num"], day).isoformat()
        except ValueError:
            fecha = ""
        for col, sede in sede_cols:
            value = ws.cell(row, col).value
            if not is_number(value):
                continue
            records.append(
                {
                    **meta,
                    "hoja": sheet_name,
                    "fila_origen": row,
                    "fecha_estimada": fecha,
                    "dia": day,
                    "fecha_original": clean_text(label),
                    "sede_original": sede,
                    "sede_normalizada": normalize_sede(sede),
                    "usuarios_activos": float(value),
                    "formula_valor": formula_at(wb_formula[sheet_name], row, col),
                }
            )
    return records


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ranking_sedes = []
    ranking_asesores = []
    usuarios_activos = []
    cobertura = []

    files = [p for p in sorted(ROOT.rglob("*.xlsx")) if not p.name.startswith("~$")]
    for path in files:
        anio, mes_num, mes_nombre, periodo = period_from_path(path)
        meta = {
            "periodo": periodo,
            "anio": anio,
            "mes_num": mes_num,
            "mes": mes_nombre,
            "fuente_archivo": str(path),
            "archivo": path.name,
        }
        wb_values = load_workbook(path, read_only=True, data_only=True)
        wb_formula = load_workbook(path, read_only=True, data_only=False)

        rank_sede_sheet = get_sheet(wb_values, "RANKING SEDE")
        rank_ases_sheet = get_sheet(wb_values, "RANKING ASESORES") or get_sheet(wb_values, "RANKING")
        activos_sheet = get_sheet(wb_values, "USUARIOS ACTIVOS") or get_sheet(wb_values, "ACTIVOS")
        dia_sheet = get_sheet(wb_values, "DIA A DIA")

        sedes = extract_rank(wb_values, wb_formula, rank_sede_sheet, "SEDE", "SEDE", meta)
        asesores = extract_rank(wb_values, wb_formula, rank_ases_sheet, "ASESOR", "ASESOR", meta)
        activos = extract_activos(wb_values, wb_formula, activos_sheet, meta)

        ranking_sedes.extend(sedes)
        ranking_asesores.extend(asesores)
        usuarios_activos.extend(activos)

        cobertura.append(
            {
                **meta,
                "hojas": len(wb_values.sheetnames),
                "tiene_dia_a_dia": bool(dia_sheet),
                "hoja_dia_a_dia": dia_sheet or "",
                "tiene_ranking_sede": bool(rank_sede_sheet),
                "registros_ranking_sede": len(sedes),
                "tiene_ranking_asesores": bool(rank_ases_sheet),
                "registros_ranking_asesores": len(asesores),
                "tiene_usuarios_activos": bool(activos_sheet),
                "registros_usuarios_activos": len(activos),
                "hojas_archivo": ", ".join(wb_values.sheetnames),
            }
        )

    data = {
        "resumen": [
            {"metrica": "Archivos Excel reales", "valor": len(files)},
            {"metrica": "Registros ranking sedes", "valor": len(ranking_sedes)},
            {"metrica": "Registros ranking asesores", "valor": len(ranking_asesores)},
            {"metrica": "Registros usuarios activos", "valor": len(usuarios_activos)},
            {"metrica": "Periodo inicial", "valor": min(r["periodo"] for r in cobertura)},
            {"metrica": "Periodo final", "valor": max(r["periodo"] for r in cobertura)},
        ],
        "cobertura": cobertura,
        "ranking_sedes": ranking_sedes,
        "ranking_asesores": ranking_asesores,
        "usuarios_activos": usuarios_activos,
    }
    OUT_FILE.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"OK {OUT_FILE}")
    for item in data["resumen"]:
        print(f"{item['metrica']}: {item['valor']}")


if __name__ == "__main__":
    main()
